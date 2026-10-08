#!/usr/bin/env python3
"""
Fantasy Life — 2027 draft research workbook generator.

    python3 draft_guide/build_workbook.py            # -> draft_guide/Fantasy_Life_2027_Draft_Research.xlsx
    python3 draft_guide/build_workbook.py --today 2026-10-20

draft_guide/data/<Category>.json is the SOURCE OF TRUTH (edit those, then rebuild); the xlsx is generated.
Every entry carries: name, detail, odds, source, as_of, confidence (A exchange/book price seen, B article
reporting odds, C estimate), plus numeric inputs:
    mu, sd                      expected outcome and its uncertainty in a category-specific unit
                                (win%, strength 0-1, composite $M, chart points, growth %, return %)
    p_champ/p_final/p_semi/...  de-vigged title ladder probabilities (team sports; inferred when missing)
    category extras             p_slam_win, p_major_win, p_top5, p_nom/p_win (Awards), grammy_ev, ...

Value model (see /fl-draft-prep): EV = E[rank points among the 13 drafted] (Monte Carlo on mu/sd vs a
plausible drafted field) + E[bonus] (ladder x probabilities; live Excel formula using the Assumptions tab).
VOR = EV - EV of the replacement-level pick (Nth best option in that category; N on the Assumptions tab).
"""
import json
import os
import sys
from datetime import date

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "data")
OUT = os.path.join(HERE, "Fantasy_Life_2027_Draft_Research.xlsx")

CATS = ["NFL", "NBA", "MLB", "NHL", "NCAAF", "NCAAB", "Tennis", "Golf", "NASCAR", "MLS",
        "Actor", "Actress", "Musician", "Country", "Stock"]
TEAM = {"NFL", "NBA", "MLB", "NHL", "NCAAF", "NCAAB", "MLS"}
# champion, lost final, lost semi/conf final, lost quarter, lost first round  (13 / 9 / 6.5 / 4 / 2.5)
LADDER = {c: [13, 9, 6.5, 4, 2.5] for c in TEAM}
LADDER["MLS"] = [13, 9, 6.5, 0, 0]
# When round odds aren't posted, infer the ladder from the title probability (equal-strength rule of
# thumb: reaching the final is ~2x as likely as winning, etc.). Tagged as inferred (confidence C).
INFER = {"p_final": 2.0, "p_semi": 4.0, "p_quarter": 8.0, "p_r1": 16.0}
P_CAP = 0.97
BONUS_CAP = 13.0
DEFAULT_REPL_RANK = 7          # "you realistically get ~the 7th-best option in a category"
SIMS = 6000

PROB_KEYS = ["p_champ", "p_final", "p_semi", "p_quarter", "p_r1"]

# Certainty layer (v2, user feedback 2026-10-08). Every category's raw mu/sd/p in data/*.json is "what the
# price/rating says today"; this layer widens it for how far away the scored event is:
#   sd_mult   multiplies sd (more variance -> flatter rank-point EV)
#   mu_shrink pulls mu toward the category's top-13 mean (year-to-year persistence is weak)
#   p_shrink  pulls title/top-5 probabilities toward the uniform field (prices that don't exist yet)
#   flat      stock: pure luck -> mu and sd forced equal for everyone (no variance premium, no favourites)
# NFL/NBA/NHL/NCAAF are in progress or opening now (and priced by Kalshi): untouched. NCAAB hasn't started.
# MLB/MLS/NASCAR score calendar 2027 with no 2027 market: heavy shrink. Golf/Tennis: predictable, untouched.
CERT = {
    "NFL": dict(), "NBA": dict(), "NHL": dict(), "NCAAF": dict(),
    "Golf": dict(), "Tennis": dict(),
    "NCAAB": dict(sd_mult=1.6, mu_shrink=0.30, p_shrink=0.25),
    "MLB": dict(sd_mult=2.0, mu_shrink=0.50, p_shrink=0.50),
    "MLS": dict(sd_mult=2.0, mu_shrink=0.50, p_shrink=0.50),
    "NASCAR": dict(sd_mult=2.0, mu_shrink=0.40, p_shrink=0.50),
    "Stock": dict(flat=True),
}
# Draft-market bias (user, 2026-10-08): the league over-drafts Grammy/Oscar names, so for these categories the
# field's pick order is modelled as baseline + w x award-EV, and the replacement level you can expect LATE is
# the best candidate outside the field's first M picks.
MARKET = {"Musician": dict(w=2.0, m=8), "Actor": dict(w=2.0, m=8), "Actress": dict(w=2.0, m=8)}

GOLD = PatternFill("solid", fgColor="FFE699")
GREEN = PatternFill("solid", fgColor="C6E0B4")
BLUE = PatternFill("solid", fgColor="DDEBF7")
GREY = PatternFill("solid", fgColor="EDEDED")
RED = PatternFill("solid", fgColor="F8CBAD")
HEAD = PatternFill("solid", fgColor="1F3864")
BOLD = Font(bold=True)
WHITE = Font(bold=True, color="FFFFFF")


def load_json(name, default=None):
    path = os.path.join(DATA, name)
    if not os.path.exists(path):
        return default
    with open(path) as f:
        return json.load(f)


def _adjust(cat, ents):
    cfg = CERT.get(cat) or {}
    if not ents or not cfg:
        return ents
    top = sorted((e.get("mu") for e in ents if e.get("mu") is not None), reverse=True)[:13]
    m = sum(top) / len(top) if top else 0.0
    n = len(ents)
    sds = sorted(e.get("sd") or 0 for e in ents)
    for e in ents:
        e["_raw"] = {"mu": e.get("mu"), "sd": e.get("sd"), "p_champ": e.get("p_champ"), "p_top5": e.get("p_top5")}
        if cfg.get("flat"):
            e["mu"], e["sd"] = round(m, 2), round(sds[len(sds) // 2], 2)
            e["confidence"] = "C"
            continue
        if e.get("mu") is not None:
            e["mu"] = round(m + (1 - cfg.get("mu_shrink", 0)) * (e["mu"] - m), 4)
        if e.get("sd") is not None:
            e["sd"] = round(e["sd"] * cfg.get("sd_mult", 1.0), 4)
        ps = cfg.get("p_shrink", 0)
        if ps and e.get("p_champ") is not None:
            e["p_champ"] = round((1 - ps) * e["p_champ"] + ps / n, 4)
            for k in PROB_KEYS[1:]:              # re-infer the ladder from the shrunk title price
                e[k] = None
        if ps and e.get("p_top5") is not None:
            e["p_top5"] = round((1 - ps) * e["p_top5"] + ps * 5 / n, 4)
    return ents


def adj_tag(cat):
    cfg = CERT.get(cat) or {}
    if cfg.get("flat"):
        return "[v2: flattened — pure luck, everyone equal]"
    if not cfg:
        return ""
    return f"[v2 certainty adj: sd x{cfg.get('sd_mult',1)}, mu shrink {int(cfg.get('mu_shrink',0)*100)}%, p shrink {int(cfg.get('p_shrink',0)*100)}%]"


def entries_for(cat):
    d = load_json(f"{cat}.json", {}) or {}
    ents = d.get("entries", [])
    for e in ents:
        e["_as_of_cat"] = d.get("as_of")
    return _adjust(cat, ents)


# ── value model ──────────────────────────────────────────────────────────────
def infer_ladder(e):
    """Fill missing ladder probabilities from p_champ; returns (probs dict, inferred flag)."""
    pT = e.get("p_champ")
    if pT is None:
        return {k: None for k in PROB_KEYS}, False
    probs = {"p_champ": pT}
    inferred = False
    prev = pT
    for k in PROB_KEYS[1:]:
        v = e.get(k)
        if v is None:
            v = min(P_CAP, INFER[k] * pT)
            inferred = True
        v = max(v, prev)            # cumulative: reaching a round is at least as likely as the next
        probs[k] = min(P_CAP, v)
        prev = probs[k]
    return probs, inferred


def bonus_ev(cat, e, assum):
    """Python mirror of the Excel bonus formula (used for sorting/tiering; Excel recomputes live)."""
    if cat in TEAM:
        probs, _ = infer_ladder(e)
        if probs["p_champ"] is None:
            return 0.0
        lad = assum["ladders"][cat]
        p = [probs[k] for k in PROB_KEYS]
        ev = lad[0] * p[0] + lad[1] * (p[1] - p[0]) + lad[2] * (p[2] - p[1]) + lad[3] * (p[3] - p[2]) + lad[4] * (p[4] - p[3])
        return min(ev, BONUS_CAP)
    if cat == "Tennis":
        return min(assum["slam_win"] * (e.get("p_slam_win") or 0) + assum["slam_ru"] * (e.get("p_slam_ru") or 0), BONUS_CAP)
    if cat == "Golf":
        return min(assum["major_win"] * (e.get("p_major_win") or 0) + assum["major_ru"] * (e.get("p_major_ru") or 0), BONUS_CAP)
    if cat == "NASCAR":
        pc, p5 = e.get("p_champ") or 0, e.get("p_top5") or 0
        return min(13 * pc + max(p5 - pc, 0) * assum["nascar_other_avg"], BONUS_CAP)
    if cat in ("Actor", "Actress"):
        return min(oscar_ev(e.get("name"), cat, assum["_awards"]), BONUS_CAP)
    if cat == "Musician":
        return min(e.get("grammy_ev") or 0, BONUS_CAP)
    return 0.0


def oscar_ev(name, cat, awards):
    """Sum over a person's contender rows for the matching gender. Lead win 13 / supp win 9 / nom 4 / 2.5."""
    gender_key = "Actor" if cat == "Actor" else "Actress"
    ev = 0.0
    for a in awards:
        if a.get("name") != name or gender_key not in (a.get("category") or ""):
            continue
        pn, pw = a.get("p_nom") or 0, a.get("p_win") or 0
        lead = "Supporting" not in a.get("category", "")
        win_pts, nom_pts = (13, 4) if lead else (9, 2.5)
        ev += win_pts * pw + nom_pts * max(pn - pw, 0)
    return ev


def oscar_detail(name, cat, awards):
    gender_key = "Actor" if cat == "Actor" else "Actress"
    out = []
    for a in awards:
        if a.get("name") == name and gender_key in (a.get("category") or ""):
            out.append(f'{a["category"]}: {a.get("film","?")} (nom {int((a.get("p_nom") or 0)*100)}%, win {int((a.get("p_win") or 0)*100)}%)')
    return "; ".join(out)


def rank_points_mc(ents, seed=2027):
    """E[rank points] and SD, simulating each entry against the best 12 OTHER entries by mu (a plausible
    drafted field). rank 1 -> 13 pts ... rank 13 -> 1 pt; ties split via continuous draws (prob. 0)."""
    n = len(ents)
    if n == 0:
        return []
    rng = np.random.default_rng(seed)
    mu = np.array([e.get("mu") if e.get("mu") is not None else 0.0 for e in ents], dtype=float)
    sd = np.array([max(e.get("sd") or 1e-6, 1e-6) for e in ents], dtype=float)
    order = np.argsort(-mu)
    out = []
    for i in range(n):
        opp = [j for j in order if j != i][:12]
        x = rng.normal(mu[i], sd[i], size=SIMS)
        opp_draw = rng.normal(mu[opp][:, None], sd[opp][:, None], size=(len(opp), SIMS))
        better = (opp_draw > x[None, :]).sum(axis=0)
        pts = 13 - np.minimum(better, 12)           # 13..1
        out.append((float(pts.mean()), float(pts.std())))
    return out


# ── workbook ─────────────────────────────────────────────────────────────────
def build(today):
    assum_file = load_json("assumptions.json", {}) or {}
    assum = {
        "ladders": {c: assum_file.get("ladders", {}).get(c, LADDER[c]) for c in TEAM},
        "slam_win": assum_file.get("slam_win", 4.0), "slam_ru": assum_file.get("slam_ru", 2.5),
        "major_win": assum_file.get("major_win", 6.0), "major_ru": assum_file.get("major_ru", 2.5),
        "nascar_other_avg": assum_file.get("nascar_other_avg", 4.5),
        "repl_rank": {c: assum_file.get("replacement_rank", {}).get(c, DEFAULT_REPL_RANK) for c in CATS},
        "risk_lambda": assum_file.get("risk_lambda", 0.0),
    }
    awards = (load_json("Awards.json", {}) or {}).get("entries", [])
    assum["_awards"] = awards

    # compute everything in Python first (sorting, tiers, QA); Excel formulas recompute the live columns
    table = {}
    for cat in CATS:
        ents = entries_for(cat)
        mc = rank_points_mc(ents)
        rows = []
        for e, (erank, sdp) in zip(ents, mc):
            b = bonus_ev(cat, e, assum)
            rows.append({"e": e, "erank": erank, "sdp": sdp, "bonus": b, "total": erank + min(b, BONUS_CAP)})
        rows.sort(key=lambda r: -r["total"])
        for i, r in enumerate(rows, 1):
            r["rank"] = i
            r["tier"] = 1 if i <= 3 else 2 if i <= 7 else 3 if i <= 13 else 4
        mk = MARKET.get(cat)
        if mk:
            for r in rows:
                r["mscore"] = r["erank"] + mk["w"] * min(r["bonus"], BONUS_CAP)
            for i, r in enumerate(sorted(rows, key=lambda r: -r["mscore"]), 1):
                r["mrank"] = i
        table[cat] = rows

    wb = Workbook()
    ws_readme = wb.active
    ws_readme.title = "README"
    ws_master = wb.create_sheet("Master")
    ws_board = wb.create_sheet("Draft Board")
    ws_as = wb.create_sheet("Assumptions")

    # Assumptions sheet (live inputs)
    ws_as.append(["Category", "Champion", "Lost final", "Lost semi", "Lost quarter", "Lost round 1",
                  "Replacement rank N", "Bonus cap", "Replacement EV (Nth best / best left after the field's picks)",
                  "Field award overweight w", "Field takes first M"])
    for c in ws_as[1]:
        c.fill, c.font = HEAD, WHITE
    as_row = {}
    for i, cat in enumerate(CATS, 2):
        as_row[cat] = i
        lad = assum["ladders"].get(cat, [None] * 5)
        ws_as.cell(i, 1, cat)
        for j, v in enumerate(lad):
            ws_as.cell(i, 2 + j, v)
        ws_as.cell(i, 7, assum["repl_rank"][cat])
        ws_as.cell(i, 8, BONUS_CAP)
        if cat in MARKET:
            ws_as.cell(i, 10, MARKET[cat]["w"])
            ws_as.cell(i, 11, MARKET[cat]["m"])
    base = len(CATS) + 4
    ws_as.cell(base, 1, "Other inputs").font = BOLD
    others = [("Tennis slam win pts", assum["slam_win"]), ("Tennis slam runner-up pts", assum["slam_ru"]),
              ("Golf major win pts", assum["major_win"]), ("Golf major runner-up pts", assum["major_ru"]),
              ("NASCAR avg pts for finishing 2nd-5th", assum["nascar_other_avg"]),
              ("Risk aversion λ (Risk-adj VOR = VOR − λ·SD of rank pts)", assum["risk_lambda"])]
    oc = {}
    for k, (label, v) in enumerate(others, base + 1):
        ws_as.cell(k, 1, label)
        ws_as.cell(k, 2, v)
        oc[label.split(" ")[0] + label.split(" ")[1]] = f"Assumptions!$B${k}"
    SLAM_WIN, SLAM_RU = f"Assumptions!$B${base+1}", f"Assumptions!$B${base+2}"
    MAJ_WIN, MAJ_RU = f"Assumptions!$B${base+3}", f"Assumptions!$B${base+4}"
    NAS_OTH, LAMBDA = f"Assumptions!$B${base+5}", f"Assumptions!$B${base+6}"
    cb = base + len(others) + 3
    ws_as.cell(cb, 1, "Certainty adjustments applied to data/*.json before the model (v2)").font = BOLD
    ws_as.cell(cb + 1, 1, "Category")
    for j, h in enumerate(["sd multiplier", "mu shrink to field", "title-prob shrink to uniform", "Why"], 2):
        ws_as.cell(cb + 1, j, h)
    why = {"NFL": "season in progress, Kalshi-priced", "NBA": "opens Oct 20, Kalshi-priced", "NHL": "season started, Kalshi-priced",
           "NCAAF": "season in progress, Kalshi-priced", "Golf": "OWGR is sticky; majors are stable at the top", "Tennis": "rankings sticky; slams stable at the top",
           "NCAAB": "season not started; tournament chaos", "MLB": "calendar 2027, no 2027 market; year-to-year win% persistence is weak",
           "MLS": "calendar 2027, no 2027 market; playoff-driven", "NASCAR": "calendar 2027, no 2027 market; Chase is a lottery",
           "Stock": "pure luck: mu and sd forced equal for all picks", "Actor": "see Market overweight", "Actress": "see Market overweight",
           "Musician": "see Market overweight", "Country": "IMF-published growth, small sd"}
    for k, cat in enumerate(CATS, cb + 2):
        cfg = CERT.get(cat) or {}
        ws_as.cell(k, 1, cat)
        ws_as.cell(k, 2, "flat" if cfg.get("flat") else cfg.get("sd_mult", 1.0))
        ws_as.cell(k, 3, "flat" if cfg.get("flat") else cfg.get("mu_shrink", 0.0))
        ws_as.cell(k, 4, "-" if cfg.get("flat") else cfg.get("p_shrink", 0.0))
        ws_as.cell(k, 5, why.get(cat, ""))
    ws_as.cell(cb + len(CATS) + 3, 1, "These are applied in Python (edit CERT in build_workbook.py and rebuild); the Excel cells above are the live ones.")
    ws_as.column_dimensions["A"].width = 58
    for col in "BCDEFGHIJK":
        ws_as.column_dimensions[col].width = 16

    # raw tabs
    headers_common_tail = ["E[rank pts] (MC)", "E[bonus]", "E[total]", "VOR", "Risk-adj VOR", "Tier",
                           "SD rank pts", "Source", "As of", "Conf", "Notes", "Field draft score", "Field rank"]
    extra_headers = {
        "team": ["p(champ)", "p(final)", "p(semi)", "p(quarter)", "p(rd 1)"],
        "Tennis": ["p(≥1 slam win)", "p(≥1 slam RU)", "Gender", "", ""],
        "Golf": ["p(≥1 major win)", "p(≥1 major RU)", "", "", ""],
        "NASCAR": ["p(champ)", "p(top 5)", "", "", ""],
        "Actor": ["Oscar p(nom)", "Oscar p(win)", "Oscar detail", "Films", ""],
        "Actress": ["Oscar p(nom)", "Oscar p(win)", "Oscar detail", "Films", ""],
        "Musician": ["Grammy EV", "p(big-4 win)", "exp. noms", "", ""],
        "Country": ["IMF 2027 fcst %", "p(top 3)", "", "", ""],
        "Stock": ["Direction", "", "", "", ""],
    }
    refs = {}      # cat -> {name: row}, plus sheet ranges for Draft Board formulas
    for cat in CATS:
        ws = wb.create_sheet(cat)
        rows = table[cat]
        tail_start = 9  # I
        ws.append([f"{cat} — 2027 draft research (EV model; see README)"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([f"As of {today}. Sorted by E[total]. Yellow = inferred / low-confidence (C)."])
        ws.append([])
        kind = "team" if cat in TEAM else cat
        hdr = ["Candidate", "Detail", "Odds / price"] + extra_headers[kind] + ["mu", "sd"] + headers_common_tail
        # columns: A name, B detail, C odds, D..H extras, I mu, J sd, K erank, L bonus, M total, N VOR, O riskadj, P tier, Q sdp, R src, S asof, T conf, U notes
        ws.append(hdr)
        for c in ws[4]:
            c.fill, c.font = HEAD, WHITE
            c.alignment = Alignment(wrap_text=True, vertical="top")
        last = 4 + len(rows)
        refs[cat] = {"first": 5, "last": last, "rows": {}}
        ar = as_row[cat]
        for i, r in enumerate(rows):
            e, n = r["e"], 5 + i
            probs, inferred = infer_ladder(e) if cat in TEAM else ({}, False)
            if cat in TEAM:
                extras = [probs.get(k) for k in PROB_KEYS]
            elif cat == "Tennis":
                extras = [e.get("p_slam_win"), e.get("p_slam_ru"), e.get("gender"), None, None]
            elif cat == "Golf":
                extras = [e.get("p_major_win"), e.get("p_major_ru"), None, None, None]
            elif cat == "NASCAR":
                extras = [e.get("p_champ"), e.get("p_top5"), None, None, None]
            elif cat in ("Actor", "Actress"):
                aw = [a for a in awards if a.get("name") == e.get("name") and ("Actor" if cat == "Actor" else "Actress") in (a.get("category") or "")]
                pn = max([a.get("p_nom") or 0 for a in aw], default=0)
                pw = max([a.get("p_win") or 0 for a in aw], default=0)
                films = "; ".join(f'{f.get("title")} ({int((f.get("p_release") or 0)*100)}%)' for f in e.get("films", []))
                extras = [pn or None, pw or None, oscar_detail(e.get("name"), cat, awards) or None, films or None, None]
            elif cat == "Musician":
                extras = [e.get("grammy_ev"), e.get("p_win_big4"), e.get("exp_noms"), None, None]
            elif cat == "Country":
                extras = [e.get("forecast_2027"), e.get("p_top3"), None, None, None]
            else:
                extras = [("Short" if "Short" in e.get("name", "") else "Long"), None, None, None, None]
            ws.append([e.get("name"), e.get("detail"), e.get("odds")] + extras + [e.get("mu"), e.get("sd")])
            # bonus formula (live) — team ladder / awards / simple
            if cat in TEAM:
                f = (f"=Assumptions!$B${ar}*D{n}+Assumptions!$C${ar}*(E{n}-D{n})+Assumptions!$D${ar}*(F{n}-E{n})"
                     f"+Assumptions!$E${ar}*(G{n}-F{n})+Assumptions!$F${ar}*(H{n}-G{n})")
                if probs.get("p_champ") is None:
                    f = 0
            elif cat == "Tennis":
                f = f"={SLAM_WIN}*N(D{n})+{SLAM_RU}*N(E{n})"
            elif cat == "Golf":
                f = f"={MAJ_WIN}*N(D{n})+{MAJ_RU}*N(E{n})"
            elif cat == "NASCAR":
                f = f"=13*N(D{n})+MAX(N(E{n})-N(D{n}),0)*{NAS_OTH}"
            else:
                f = round(r["bonus"], 3)         # Oscar/Grammy EV computed in Python from Awards.json / entry
            ws.cell(n, 11, round(r["erank"], 3))
            ws.cell(n, 12, f)
            ws.cell(n, 13, f"=K{n}+MIN(L{n},Assumptions!$H${ar})")
            ws.cell(n, 14, f"=M{n}-Assumptions!$I${ar}")
            ws.cell(n, 15, f"=N{n}-{LAMBDA}*Q{n}")
            ws.cell(n, 16, r["tier"])
            ws.cell(n, 17, round(r["sdp"], 3))
            ws.cell(n, 18, e.get("source"))
            ws.cell(n, 19, e.get("as_of") or e.get("_as_of_cat"))
            ws.cell(n, 20, e.get("confidence"))
            ws.cell(n, 21, ((e.get("notes") or "") + " " + adj_tag(cat)).strip())
            if cat in MARKET:
                ws.cell(n, 22, f"=K{n}+Assumptions!$J${ar}*MIN(L{n},Assumptions!$H${ar})")
                ws.cell(n, 23, f"=RANK(V{n},$V$5:$V${last})")
            fill = GOLD if r["tier"] == 1 else GREEN if r["tier"] == 2 else BLUE if r["tier"] == 3 else None
            if fill:
                ws.cell(n, 1).fill = fill
            if e.get("confidence") == "C" or inferred:
                ws.cell(n, 20).fill = GOLD
            refs[cat]["rows"][e.get("name")] = n
        ws.cell(last + 2, 1, "Replacement-level EV (Nth best)").font = BOLD
        nth = f"LARGE('{cat}'!$M$5:$M${last},MIN(G{ar},COUNT('{cat}'!$M$5:$M${last})))"
        if cat in MARKET:       # best candidate the field leaves after its first M award-driven picks
            left = f"SUMPRODUCT(MAX(('{cat}'!$W$5:$W${last}>K{ar})*'{cat}'!$M$5:$M${last}))"
            ws_as.cell(ar, 9, f"=MAX({nth},{left})")
        else:
            ws_as.cell(ar, 9, f"={nth}")
        ws.freeze_panes = "B5"
        for col, w in zip("ABCDEFGHIJKLMNOPQRSTUVW", [30, 36, 22, 11, 11, 11, 11, 11, 9, 9, 11, 10, 10, 9, 11, 6, 9, 38, 11, 6, 60, 12, 9]):
            ws.column_dimensions[col].width = w

    # Draft Board: every candidate, sorted by default-N VOR
    repl_ev = {}
    for cat in CATS:
        tots = sorted([r["total"] for r in table[cat]], reverse=True)
        n = min(assum["repl_rank"][cat], len(tots)) if tots else 0
        repl_ev[cat] = tots[n - 1] if n else 0.0
        if cat in MARKET:
            left = [r["total"] for r in table[cat] if r["mrank"] > MARKET[cat]["m"]]
            repl_ev[cat] = max(repl_ev[cat], max(left, default=0.0))
    board = []
    for cat in CATS:
        for r in table[cat]:
            board.append((r["total"] - repl_ev[cat], cat, r))
    board.sort(key=lambda t: -t[0])

    def plan(cat, r):
        if cat in MARKET:
            hot = r["mrank"] <= MARKET[cat]["m"]
            if hot and r["rank"] <= 7:
                return "Good but the field reaches for it: let it come to you"
            if hot:
                return "Overpriced by the field on award buzz: skip"
            if r["rank"] <= 7:
                return "TARGET LATE: value the field undervalues"
            return "late filler"
        if cat == "Stock":
            return "coin flip: take last, no edge"
        if cat in ("MLB", "MLS", "NASCAR", "NCAAB"):
            return "low confidence: far-out event"
        return ""
    ws_board.append(["Rank", "Category", "Candidate", "E[total]", "VOR", "Risk-adj VOR", "Tier in category", "Odds / price", "Detail", "Conf", "Source", "As of", "Plan"])
    for c in ws_board[1]:
        c.fill, c.font = HEAD, WHITE
    for i, (vor, cat, r) in enumerate(board, 1):
        n = refs[cat]["rows"][r["e"].get("name")]
        q = f"'{cat}'"
        ws_board.append([i, cat, r["e"].get("name"), f"={q}!M{n}", f"={q}!N{n}", f"={q}!O{n}", r["tier"],
                         r["e"].get("odds"), r["e"].get("detail"), r["e"].get("confidence"), r["e"].get("source"),
                         r["e"].get("as_of") or r["e"].get("_as_of_cat"), plan(cat, r)])
        if i <= 15:
            ws_board.cell(i + 1, 3).fill = GOLD
    ws_board.freeze_panes = "D2"
    ws_board.auto_filter.ref = f"A1:M{len(board)+1}"
    for col, w in zip("ABCDEFGHIJKLM", [6, 11, 30, 10, 9, 12, 9, 22, 40, 6, 40, 11, 44]):
        ws_board.column_dimensions[col].width = w

    # Master (same shape as the 2025 workbook: one column per category, tier-ordered)
    ws_master.append(["Top 15 (EV) →"] + CATS)
    for c in ws_master[1]:
        c.fill, c.font = HEAD, WHITE
    for i in range(15):
        row = [f"#{i+1}"]
        for cat in CATS:
            rows = table[cat]
            if i < len(rows):
                r = rows[i]
                odds = f" {r['e'].get('odds')}" if r["e"].get("odds") else ""
                row.append(f"{r['e'].get('name')}{odds} · EV {r['total']:.1f}")
            else:
                row.append("")
        ws_master.append(row)
        for j in range(len(CATS)):
            cell = ws_master.cell(i + 2, j + 2)
            cell.fill = GOLD if i < 3 else GREEN if i < 7 else BLUE if i < 13 else GREY
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws_master.column_dimensions["A"].width = 14
    for j in range(len(CATS)):
        ws_master.column_dimensions[get_column_letter(j + 2)].width = 30
    ws_master.freeze_panes = "B2"

    # QA
    ws_qa = wb.create_sheet("QA")
    ws_qa.append(["Severity", "Category", "Candidate", "Issue"])
    for c in ws_qa[1]:
        c.fill, c.font = HEAD, WHITE
    top10 = {(cat, r["e"].get("name")) for _, cat, r in board[:10]}
    for cat in CATS:
        names = set()
        for r in table[cat]:
            e, nm = r["e"], r["e"].get("name")
            if nm in names:
                ws_qa.append(["WARN", cat, nm, "duplicate candidate"])
            names.add(nm)
            if not e.get("source"):
                ws_qa.append(["ERROR", cat, nm, "no source"])
            asof = e.get("as_of") or e.get("_as_of_cat")
            if not asof:
                ws_qa.append(["ERROR", cat, nm, "no as-of date"])
            else:
                try:
                    age = (today - date.fromisoformat(asof)).days
                    if age > 7:
                        ws_qa.append(["WARN", cat, nm, f"stale: {age} days old"])
                except ValueError:
                    ws_qa.append(["WARN", cat, nm, f"bad as-of date {asof!r}"])
            if e.get("mu") is None or e.get("sd") is None:
                ws_qa.append(["WARN", cat, nm, "missing mu/sd — baseline EV is a placeholder"])
            if e.get("confidence") == "C" and (cat, nm) in top10:
                ws_qa.append(["WARN", cat, nm, "confidence-C candidate is in the overall board top 10 — verify before relying"])
        if cat in TEAM and cat not in ("NCAAF", "NCAAB"):
            tot = sum((r["e"].get("p_champ") or 0) for r in table[cat])
            if table[cat] and not 0.8 <= tot <= 1.2:
                ws_qa.append(["WARN", cat, "", f"title probabilities sum to {tot:.2f} (expect ~1.0 over the full field)"])
        if len(table[cat]) < 13:
            ws_qa.append(["WARN", cat, "", f"only {len(table[cat])} candidates (need >= 13 for a drafted field)"])
    for col, w in zip("ABCD", [10, 12, 32, 80]):
        ws_qa.column_dimensions[col].width = w

    # Aliases
    ws_al = wb.create_sheet("Aliases")
    ws_al.append(["Alias (what you might type / sources use)", "Canonical (what the board uses)"])
    for c in ws_al[1]:
        c.fill, c.font = HEAD, WHITE
    for a, b in (load_json("aliases.json", {}) or {}).items():
        ws_al.append([a, b])
    ws_al.column_dimensions["A"].width = 44
    ws_al.column_dimensions["B"].width = 44

    # README
    lines = [
        ("Fantasy Life 2027 — Draft Research", True),
        (f"Built {today} for the Oct 25, 2026 draft. Regenerate: python3 draft_guide/build_workbook.py (data in draft_guide/data/*.json).", False),
        ("", False),
        ("How to read it", True),
        ("Master = top-15 per category by expected league points (EV). Draft Board = every candidate across all 15 categories ranked by VOR.", False),
        ("EV = E[rank points among the 13 drafted picks] (Monte Carlo on mu/sd) + E[bonus] (ladder × probabilities; live formula on the Assumptions tab).", False),
        ("VOR = EV − EV of the replacement pick (Nth best in the category; N on Assumptions). Risk-adj VOR subtracts λ × SD of rank points.", False),
        ("Yellow confidence cell = C (estimate) or ladder probabilities inferred from the title price. A = exchange/book price seen; B = article reporting odds.", False),
        ("", False),
        ("Timeline (league year 2027)", True),
        ("NFL 2026 and NCAAF 2026 are in progress (current records matter); NBA/NHL/NCAAB 2026-27 are starting; MLB/MLS/NASCAR/Golf/Tennis/Stock/Actor/Actress/Musician/Country are calendar 2027.", False),
        ("Oscars (Mar 2027) honour 2026 films and Grammys (Feb 2027) honour Aug 2025–Aug 2026 music → awards bonuses are priced from awards-season odds; Actor/Actress BASELINE uses 2027 releases.", False),
        ("IMF Oct 2026 WEO publishes Oct 13 — Country is provisional until then.", False),
        ("", False),
        ("Assumptions to confirm", True),
        ("Country 2027 = GDP growth only (no Olympics/World Cup in 2027). Slot-agnostic board (replacement rank is a single tunable N per category). Each player picks one per category.", False),
        ("Known limits: no 2027 futures are posted yet for MLB/MLS/NASCAR/Golf/Tennis; those rows are model estimates (conf C). Rank-point EV assumes the field drafts the best candidates by mu.", False),
        ("", False),
        ("How much to trust each tab (v2, Oct 8 2026)", True),
        ("SOLID (exchange prices, conf A): NFL, NBA, NHL, NCAAF title + conference-champion odds (Kalshi); Oscar win/nomination prices; Grammy AOTY/ROTY/SOTY/Best New Artist prices. These events are in progress or start within weeks, so they keep full weight.", False),
        ("PREDICTABLE: Golf (OWGR) and Tennis (ATP/WTA rankings) — the top is sticky year to year; probabilities are still my model, not posted odds.", False),
        ("DISCOUNTED FOR DISTANCE (v2): MLB, MLS, NASCAR score calendar 2027 with no 2027 market, so sd x2, mu pulled 40-50% toward the field, title probabilities pulled 50% toward uniform. NCAAB (not started) sd x1.6, mu 30%, probs 25%. See Assumptions → 'Certainty adjustments'.", False),
        ("STOCK = COIN FLIP: mu and sd forced equal for every pick (no favourites, no variance premium). The tab is a list, not a ranking; take it last.", False),
        ("COUNTRY: IMF-published growth is the mu (see the Country tab for the vintage used); small sd for stable economies, wide only for commodity/fragile states.", False),
        ("SOFT: Actor/Actress baseline (Wikipedia + Box Office Mojo billed casts x my box-office/RT comps — cameos not captured; Secret Wars may slip to 2028); Musician baseline (no 2027 album dates found).", False),
        ("", False),
        ("Draft-market bias (the edge)", True),
        ("The league over-drafts Grammy/Oscar names. For Musician, Actor and Actress the 'Field draft score' column (V) = baseline + w x award-EV (w on Assumptions) predicts the order your leaguemates pick in; the replacement level used for VOR is the best candidate OUTSIDE the field's first M picks.", False),
        ("Draft Board 'Plan' column: TARGET LATE = high EV but the field ignores it (box-office draws); overpriced = the field takes it on award buzz; let it come to you = good but the field reaches. For Actor/Actress the unexploited edge is real box-office draws.", False),
        ("Reading the board: NBA co-favourites and the Kalshi-priced NFL/NHL/NCAAF favourites carry the safest edge; Musician/Actor/Actress show big raw EV but are meant to be taken LATE.", False),
        ("Refresh: see .claude/commands/fl-draft-prep.md — new prices go in draft_guide/raw/*.json, run kalshi_overlay.py, edit data/*.json, rebuild.", False),
    ]
    for t, b in lines:
        ws_readme.append([t])
        if b:
            ws_readme.cell(ws_readme.max_row, 1).font = Font(bold=True, size=12)
    ws_readme.column_dimensions["A"].width = 160

    wb.save(OUT)
    return table, board


def main(argv):
    today = date.today()
    if "--today" in argv:
        today = date.fromisoformat(argv[argv.index("--today") + 1])
    table, board = build(today)
    print(f"wrote {OUT}")
    for cat in CATS:
        rows = table[cat]
        print(f"  {cat:9} {len(rows):3} candidates" + (f"  top: {rows[0]['e'].get('name')} (EV {rows[0]['total']:.1f})" if rows else ""))
    print("board top 15:")
    for i, (vor, cat, r) in enumerate(board[:15], 1):
        print(f"  {i:2}. {cat:8} {r['e'].get('name'):32} EV {r['total']:5.1f}  VOR {vor:5.1f}")


if __name__ == "__main__":
    main(sys.argv[1:])
